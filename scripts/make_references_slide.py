"""Append a one-slide reference list to the project deck.

The deck is a free-form design with no title placeholders, so this rebuilds its visual
language explicitly rather than relying on a layout: the grey eyebrow pill at top left,
a 24pt Montserrat title, and body text in Source Sans 3 at the deck's 0.59" margin and
8.82" content width. Those values were read off the existing slides, not guessed — a
slide sitting half an inch out of alignment looks worse than no slide.

``REFERENCES`` holds all 36 sources from
``docs/paper/literature-review-and-methodology.md`` with the report's numbering.
``SELECTION`` names the subset the slide shows, in display order, and those are
renumbered 1..N so the slide reads as a self-contained list. The deck and the report
therefore use different numbers for the same paper — deliberate, because a slide with
gaps like [1] [2] [5] [11] [27] looks like a mistake, and a deck is read on its own.

The deck is copied to a timestamped backup before it is touched.

Usage::

    python scripts/make_references_slide.py --deck "<path to .pptx>"
"""

from __future__ import annotations

import argparse
import copy
import re
import shutil
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

# Measured from the deck's own slides, not chosen.
MARGIN = 0.59
CONTENT_W = 8.82
BODY_FONT = "Source Sans 3"
TITLE_FONT = "Montserrat Bold"
INK = RGBColor(0x3D, 0x38, 0x38)
BLACK = RGBColor(0x00, 0x00, 0x00)

# The eyebrow pill is always "Shape 0", but the bottom-right badge is numbered per
# slide ("Rectangle: Rounded Corners 12/16/22"), so it is matched by prefix.
PILL_SHAPE = "Shape 0"
BADGE_PREFIX = "Rectangle: Rounded Corners"

# Slide 3's header geometry: the canonical content-slide layout in this deck. The pill
# sits at a different height on other slides, so it is repositioned after copying rather
# than inherited from whichever slide happened to be last.
PILL_TOP = 1.34
PILL_H = 0.22
EYEBROW_TOP = 1.385
TITLE_TOP = 1.63
SUBTITLE_TOP = 2.14
COLUMNS_TOP = 2.46
COLUMNS_BOTTOM = 7.02

# The subset shown on the slide, by report number, in display order: representation
# and detection foundations, line-level localization (the closest prior art), hybrid
# fusion (the comparators for the novelty claim), explainability, then the datasets we
# actually consume and the label-quality work that bounds every number we report.
SELECTION = (1, 2, 3, 7, 5, 6, 10, 11, 13, 18, 25, 27, 28, 29, 31)

REFERENCES = [
    "[1] Yamaguchi et al., Modeling and Discovering Vulnerabilities with Code Property "
    "Graphs. IEEE S&P, 2014.",
    "[2] Zhou et al., Devign: Effective Vulnerability Identification via Graph Neural "
    "Networks. NeurIPS, 2019.",
    "[3] Feng et al., CodeBERT: A Pre-Trained Model for Programming and Natural "
    "Languages. EMNLP Findings, 2020.",
    "[4] Guo et al., GraphCodeBERT: Pre-training Code Representations with Data Flow. "
    "ICLR, 2021.",
    "[5] Fu & Tantithamthavorn, LineVul: A Transformer-based Line-Level Vulnerability "
    "Prediction. MSR, 2022.",
    "[6] Hin et al., LineVD: Statement-level Vulnerability Detection using Graph Neural "
    "Networks. MSR, 2022.",
    "[7] Li, Wang & Nguyen, Vulnerability Detection with Fine-grained Interpretations. "
    "ESEC/FSE, 2021.",
    "[8] Learning to Quantize Vulnerability Patterns and Match to Locate "
    "Statement-Level Vulnerabilities. arXiv:2306.06109, 2023.",
    "[9] If LLMs Would Just Look: Simple Line-by-line Checking Improves Vulnerability "
    "Localization. arXiv:2410.15288, 2024.",
    "[10] Evaluating line-level localization ability of learning-based code "
    "vulnerability detection models. Machine Learning, 2025.",
    "[11] Vul-LMGNNs: Fusing Language Models and Online-Distilled GNNs for Code "
    "Vulnerability Detection. arXiv:2404.14719, 2024.",
    "[12] Transformer-based statement level vulnerability detection by cross-modal "
    "fine-grained features capture. Knowledge-Based Systems, 2025.",
    "[13] DCVD: Dual-Channel Cross-Modal Fusion for Joint Vulnerability Detection and "
    "Localization. arXiv:2605.11015, 2026.",
    "[14] Enhancing vulnerability detection by fusing code semantic features with "
    "LLM-generated explanations. Information Fusion, 2025.",
    "[15] VulReaD: Knowledge-Graph-guided Software Vulnerability Reasoning and "
    "Detection. arXiv:2602.10787, 2026.",
    "[16] An Unbiased Transformer Source Code Learning with Semantic Vulnerability "
    "Graph. arXiv:2304.11072, 2023.",
    "[17] Towards Causal Deep Learning for Vulnerability Detection. arXiv:2310.07958, " "2023.",
    "[18] Towards Explainable Vulnerability Detection with Large Language Models. "
    "arXiv:2406.09701, 2024.",
    "[19] Large Language Models in Software Security: A Survey of Vulnerability "
    "Detection Techniques. arXiv:2502.07049, 2025.",
    "[20] Everything You Wanted to Know About LLM-based Vulnerability Detection But "
    "Were Afraid to Ask. arXiv:2504.13474, 2025.",
    "[21] Revisiting Pre-trained Language Models for Vulnerability Detection. "
    "arXiv:2507.16887, 2025.",
    "[22] VULPO: Context-Aware Vulnerability Detection via On-Policy LLM Optimization. "
    "arXiv:2511.11896, 2025.",
    "[23] Graph Neural Networks for Software Vulnerability Mining: A Review. " "Information, 2026.",
    "[24] Ying et al., GNNExplainer: Generating Explanations for Graph Neural Networks. "
    "NeurIPS, 2019.",
    "[25] Jain & Wallace, Attention is not Explanation. NAACL-HLT, 2019.",
    "[26] DeYoung et al., ERASER: A Benchmark to Evaluate Rationalized NLP Models. ACL, " "2020.",
    "[27] Fan et al., A C/C++ Code Vulnerability Dataset with Code Changes and CVE "
    "Summaries. MSR, 2020.",
    "[28] Croft, Babar & Kholoosi, Data Quality for Software Vulnerability Datasets. "
    "ICSE, 2023.",
    "[29] Ding et al., Vulnerability Detection with Code Language Models: How Far Are "
    "We? ICSE, 2025.",
    "[30] CleanVul: Automatic Function-Level Vulnerability Detection in Code Commits "
    "Using LLM-Based Noise Reduction. arXiv:2411.17274, 2024.",
    "[31] Chen et al., DiverseVul: A New Vulnerable Source Code Dataset for Deep "
    "Learning Based Vulnerability Detection. RAID, 2023.",
    "[32] Ni et al., MegaVul: A C/C++ Vulnerability Dataset with Comprehensive Code "
    "Representations. MSR, 2024.",
    "[33] ActiveClean: Generating Line-Level Vulnerability Data via Active Learning. "
    "arXiv:2312.01588, 2023.",
    "[34] Bhandari, Naseer & Moonen, CVEfixes: Automated Collection of Vulnerabilities "
    "and Their Fixes from Open-Source Software. PROMISE, 2021.",
    "[35] NIST Software Assurance Reference Dataset, Juliet Test Suite for C/C++, " "v1.3.1.",
    "[36] Brunsfeld et al., Tree-sitter: An incremental parsing system for programming " "tools.",
]


def select(numbers: tuple[int, ...]) -> list[str]:
    """Return the chosen references, renumbered from 1 in display order.

    Args:
        numbers: Report reference numbers, in the order they should appear.

    Returns:
        Entry strings prefixed ``[1]``, ``[2]``, ...
    """
    out = []
    for position, number in enumerate(numbers, start=1):
        body = re.sub(r"^\[\d+\]\s*", "", REFERENCES[number - 1])
        out.append(f"[{position}] {body}")
    return out


def span(numbers: tuple[int, ...]) -> str:
    """Return the publication-year range of the selection.

    Args:
        numbers: Report reference numbers.

    Returns:
        A string like ``2014-2026``.
    """
    # Anchored to the trailing ", YYYY." so an arXiv id is never misread as a year.
    years = [
        int(match.group(1))
        for number in numbers
        if (match := re.search(r",\s*((?:19|20)\d{2})\.\s*$", REFERENCES[number - 1]))
    ]
    return f"{min(years)}-{max(years)}" if years else ""


def add_textbox(slide, left: float, top: float, width: float, height: float):
    """Add a transparent, tight-margined text box.

    Args:
        slide: Target slide.
        left: Left offset, inches.
        top: Top offset, inches.
        width: Width, inches.
        height: Height, inches.

    Returns:
        The new shape's text frame.
    """
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    return frame


def style(run, *, size: float, font: str, color: RGBColor, bold: bool = False) -> None:
    """Apply the deck's typography to one run.

    Args:
        run: Run to style.
        size: Point size.
        font: Typeface name.
        color: Text colour.
        bold: Whether to embolden.
    """
    run.font.size = Pt(size)
    run.font.name = font
    run.font.color.rgb = color
    run.font.bold = bold


def build(deck_path: Path, columns: int, font_size: float) -> None:
    """Append the references slide to the deck, after backing it up.

    Args:
        deck_path: The .pptx to modify in place.
        columns: Number of columns to flow references into.
        font_size: Point size for reference entries.
    """
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = deck_path.with_name(f"{deck_path.stem}.backup-{stamp}.pptx")
    shutil.copy2(deck_path, backup)
    print(f"backup written: {backup.name}")

    prs = Presentation(str(deck_path))
    donor = prs.slides[len(prs.slides) - 1]
    slide = prs.slides.add_slide(donor.slide_layout)

    # A freshly added slide inherits the layout's placeholders; this design uses none.
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)

    # Reuse the real decorative shapes so fill and corner radius match exactly.
    for shape in donor.shapes:
        if shape.name == PILL_SHAPE or shape.name.startswith(BADGE_PREFIX):
            clone = copy.deepcopy(shape._element)
            slide.shapes._spTree.append(clone)

    for shape in slide.shapes:
        if shape.name == PILL_SHAPE:
            shape.left, shape.top = Inches(MARGIN), Inches(PILL_TOP)
            shape.width, shape.height = Inches(0.85), Inches(PILL_H)

    frame = add_textbox(slide, MARGIN + 0.09, EYEBROW_TOP, 3.0, 0.14)
    run = frame.paragraphs[0].add_run()
    run.text = "REFERENCES"
    style(run, size=8.5, font=BODY_FONT, color=INK)

    frame = add_textbox(slide, MARGIN, TITLE_TOP, CONTENT_W, 0.42)
    run = frame.paragraphs[0].add_run()
    run.text = "References"
    style(run, size=24, font=TITLE_FONT, color=BLACK, bold=True)

    frame = add_textbox(slide, MARGIN, SUBTITLE_TOP, CONTENT_W, 0.20)
    run = frame.paragraphs[0].add_run()
    run.text = (
        f"{len(SELECTION)} key sources spanning {span(SELECTION)}. "
        f"The full {len(REFERENCES)}-reference list is in the report."
    )
    style(run, size=9, font=BODY_FONT, color=INK)

    entries = select(SELECTION)

    gap = 0.14
    col_w = (CONTENT_W - gap * (columns - 1)) / columns
    per_col = -(-len(entries) // columns)
    top = COLUMNS_TOP
    height = COLUMNS_BOTTOM - top

    for index in range(columns):
        block = entries[index * per_col : (index + 1) * per_col]
        if not block:
            continue
        frame = add_textbox(slide, MARGIN + index * (col_w + gap), top, col_w, height)
        for position, entry in enumerate(block):
            para = frame.paragraphs[0] if position == 0 else frame.add_paragraph()
            para.space_after = Pt(5.0)
            para.line_spacing = 0.92
            run = para.add_run()
            run.text = entry
            style(run, size=font_size, font=BODY_FONT, color=INK)

    prs.save(str(deck_path))
    print(f"appended slide {len(prs.slides)} to {deck_path.name}")


def main() -> None:
    """Parse arguments and build the slide."""
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--deck", required=True, help="Path to the .pptx to append to.")
    ap.add_argument("--columns", type=int, default=2)
    ap.add_argument("--font-size", type=float, default=11.0)
    args = ap.parse_args()

    deck = Path(args.deck)
    if not deck.exists():
        raise SystemExit(f"Deck not found: {deck}")
    build(deck, args.columns, args.font_size)


if __name__ == "__main__":
    main()
